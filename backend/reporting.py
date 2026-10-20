from collections import Counter
from io import BytesIO

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.history import (
    get_observations,
    get_statistics,
    time_window,
)

router = APIRouter(prefix="/api", tags=["History and reports"])
REPORT_ROW_LIMIT = 50000


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(
        ("=", "+", "-", "@")
    ):
        return "'" + value
    return value


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)

    for row in rows:
        sheet.append([safe_cell(value) for value in row])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="16324F",
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )
        sheet.column_dimensions[column[0].column_letter].width = min(
            max(length + 2, 14), 60
        )


@router.get("/observations")
def observations(
    days: int = Query(7, ge=1, le=90),
    limit: int = Query(1000, ge=1, le=5000),
):
    start, end = time_window(days)
    return {
        "start_utc": start,
        "end_utc": end,
        "statistics": get_statistics(start, end),
        "items": get_observations(start, end, limit),
        "limit": limit,
    }


@router.get("/reports/weekly.xlsx")
def weekly_report(days: int = Query(7, ge=1, le=90)):
    start, end = time_window(days)
    statistics = get_statistics(start, end)

    if statistics["observation_count"] > REPORT_ROW_LIMIT:
        raise HTTPException(
            status_code=413,
            detail="Too many rows. Choose a shorter reporting period.",
        )

    records = get_observations(start, end, REPORT_ROW_LIMIT)
    workbook = Workbook()
    workbook.remove(workbook.active)

    add_sheet(
        workbook,
        "Summary",
        ["Metric", "Value"],
        [
            ["Reporting start UTC", start],
            ["Reporting end UTC (exclusive)", end],
            ["Observations", statistics["observation_count"]],
            ["Unique IPs", statistics["unique_ips"]],
            ["Collections", statistics["collection_count"]],
            ["Countries", statistics["country_count"]],
            ["Source", "AbuseIPDB blacklist observations"],
            ["Timestamp meaning", "Collection time, not attack time"],
        ],
    )

    columns = [
        "observed_at", "collection_id", "source", "ip",
        "country", "region", "city", "latitude", "longitude",
        "asn", "organization", "abuse_confidence",
    ]

    add_sheet(
        workbook,
        "Observations",
        columns,
        [[record.get(column) for column in columns] for record in records],
    )

    indicators = {}

    for record in records:
        ip = record["ip"]
        entry = indicators.setdefault(ip, {
            "first_seen": record["observed_at"],
            "last_seen": record["observed_at"],
            "count": 0,
            "scores": [],
        })
        entry["first_seen"] = min(
            entry["first_seen"], record["observed_at"]
        )
        entry["last_seen"] = max(
            entry["last_seen"], record["observed_at"]
        )
        entry["count"] += 1
        if record["abuse_confidence"] is not None:
            entry["scores"].append(record["abuse_confidence"])

    add_sheet(
        workbook,
        "Unique indicators",
        [
            "IP", "First observed UTC", "Last observed UTC",
            "Observation count", "Maximum reputation score",
        ],
        [
            [
                ip, entry["first_seen"], entry["last_seen"],
                entry["count"],
                max(entry["scores"]) if entry["scores"] else None,
            ]
            for ip, entry in sorted(indicators.items())
        ],
    )

    geographic_counts = Counter(
        (
            record["country"] or "Unknown",
            record["region"] or "Unknown",
            record["city"] or "Unknown",
        )
        for record in records
    )

    add_sheet(
        workbook,
        "Geography",
        ["Country", "Region", "City", "Observation count"],
        [
            [*location, count]
            for location, count in geographic_counts.most_common()
        ],
    )

    asn_counts = Counter(
        record["asn"] or "Unknown" for record in records
    )

    add_sheet(
        workbook,
        "Networks",
        ["ASN", "Observation count"],
        [[asn, count] for asn, count in asn_counts.most_common()],
    )

    add_sheet(
        workbook,
        "Limitations",
        ["Topic", "Explanation"],
        [
            ["Scope", "IP reputation observations, not confirmed DDoS events."],
            ["Counts", "Repeated observations are not separate attacks."],
            ["Time", "No historical timestamps were inferred from the cache."],
            ["Geography", "IP locations are approximate provider results."],
            ["Networks", "ASN association does not prove compromise."],
            ["Coverage", "Missing observations do not establish safety."],
            ["ML", "Existing model is not used to assert DDoS detection."],
        ],
    )

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"intelligence-report-{end[:10]}.xlsx"

    return Response(
        content=buffer.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )
